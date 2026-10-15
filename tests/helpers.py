"""Fixture builders: generate PDF / DOCX / MD documents in memory."""
from __future__ import annotations

import io
from typing import List, Tuple


def make_md(items: List[Tuple[str, str]]) -> bytes:
    """items: (kind, text) with kind in h1/h2/p/li."""
    out = []
    for kind, text in items:
        if kind.startswith("h"):
            out.append("#" * int(kind[1]) + " " + text)
        elif kind == "li":
            out.append("- " + text)
        else:
            out.append(text)
        out.append("")
    return "\n".join(out).encode("utf-8")


def make_docx(items: List[Tuple[str, str]]) -> bytes:
    from docx import Document

    d = Document()
    for kind, text in items:
        if kind.startswith("h"):
            d.add_heading(text, level=int(kind[1]))
        elif kind == "li":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "table":
            rows = [r.split("|") for r in text.split(";")]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, r in enumerate(rows):
                for j, cell in enumerate(r):
                    t.cell(i, j).text = cell
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(pages: List[List[Tuple[str, str]]]) -> bytes:
    """pages: list of pages; each page list of (kind, text); kind h1 or p. ASCII text recommended."""
    import fitz

    doc = fitz.open()
    for items in pages:
        page = doc.new_page()
        y = 60
        for kind, text in items:
            size = 18 if kind.startswith("h") else 11
            font = "hebo" if kind.startswith("h") else "helv"
            rect = fitz.Rect(50, y, 545, y + 60)
            page.insert_textbox(rect, text, fontsize=size, fontname=font)
            y += 72
        page.insert_textbox(fitz.Rect(250, 810, 350, 835), str(doc.page_count), fontsize=9)
    return doc.tobytes()


def long_doc_items(n_articles: int = 40, change_article: int = None, new_text: str = None):
    items = []
    for a in range(1, n_articles + 1):
        items.append(("h2", "Dieu %d. Quy dinh so %d" % (a, a)))
        for k in range(1, 4):
            text = "Khoan %d cua Dieu %d quy dinh ve nghia vu thuc hien so %d voi noi dung rieng biet %d." % (
                k, a, a * 10 + k, a * 7 + k)
            if change_article == a and k == 2 and new_text:
                text = new_text
            items.append(("p", text))
    return items


def paginate(items, per_page: int = 10):
    return [items[i:i + per_page] for i in range(0, len(items), per_page)]
